import io
from docx import Document
from sqlalchemy import func, select
from backend.app.models import Projeto


def test_preview_preserva_incompletos_sem_gravar_projeto(client, factory):
    doc = Document()
    doc.add_paragraph('PSC 3000')
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ['Destino','Equipamento','Origem','Responsável','Ação']):
        cell.text = text
    for values in [['SITE A','2 Módulos','Engenharia','Implantação','Adquirir'],
                   ['SITE B','Cordão sem quantidade','Engenharia','Implantação','Adquirir']]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    data = io.BytesIO(); doc.save(data)
    response = client.post('/documentos/extrair', files={'arquivo':('PSC 3000.docx',data.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['codigo'] == 'PSC 3000' and len(result['itens']) == 2
    assert result['itens'][1]['quantidade'] is None and result['itens'][1]['revisao']
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(Projeto)) == 0


def test_preview_rejeita_formato_invalido_e_vazio(client):
    assert client.post('/documentos/extrair', files={'arquivo':('arquivo.exe',b'abc')}).status_code == 422
    assert client.post('/documentos/extrair', files={'arquivo':('arquivo.pdf',b'')}).status_code == 422
