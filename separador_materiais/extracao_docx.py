"""
extracao_docx.py
================

Leitura da tabela "Lista de equipamentos / Material utilizado" de arquivos
**DOCX** e **DOC**.

- DOCX: lido direto com ``python-docx`` (a tabela costuma ser uma tabela real
  do Word). Não precisa do Microsoft Word instalado.
- DOC (formato antigo): convertido para DOCX via automação do **Microsoft Word**
  (win32com) e então lido como DOCX. Requer o Word instalado.

A saída é um ``ResultadoExtracao`` com DataFrame nas colunas do cabeçalho
esperado, pronto para ``normalizacao.normalizar_tabela``.
"""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Optional

import pandas as pd

from extracao import ErroExtracao, ResultadoExtracao
from normalizacao import CABECALHO_ESPERADO, normalizar_texto

_CAMPOS = {normalizar_texto(c) for c in CABECALHO_ESPERADO}
_ORIG_POR_NORM = {normalizar_texto(c): c for c in CABECALHO_ESPERADO}

_PADROES_PSC = [
    ("PSC", re.compile(r"PSC[\s:_/#.\-]*(\d{2,6})", re.IGNORECASE)),
    ("PS", re.compile(r"\bPS[\s:_/#.\-]*(\d{2,6})", re.IGNORECASE)),
]


def _detectar_psc(texto: str) -> Optional[str]:
    for prefixo, padrao in _PADROES_PSC:
        m = padrao.search(texto or "")
        if m:
            return f"{prefixo} {m.group(1)}"
    return None


def _cabecalho_casa(celulas) -> bool:
    norms = {normalizar_texto(c) for c in celulas}
    return _CAMPOS.issubset(norms)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def extrair_tabela_docx(conteudo: bytes, nome_arquivo: str) -> ResultadoExtracao:
    """Extrai a tabela de materiais de um DOCX (tabela real do Word)."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ErroExtracao("Biblioteca python-docx não instalada (necessária p/ DOCX).") from exc

    try:
        doc = Document(io.BytesIO(conteudo))
    except Exception as exc:  # noqa: BLE001
        raise ErroExtracao(f"Não foi possível abrir o DOCX (arquivo inválido?): {exc}") from exc

    for tabela in doc.tables:
        if not tabela.rows:
            continue
        cabecalho = [c.text for c in tabela.rows[0].cells]
        if not _cabecalho_casa(cabecalho):
            continue

        # mapeia índice de coluna -> campo interno (1ª ocorrência de cada)
        idx_por_campo: dict[str, int] = {}
        for j, h in enumerate(cabecalho):
            n = normalizar_texto(h)
            if n in _CAMPOS and n not in idx_por_campo:
                idx_por_campo[n] = j

        linhas = []
        for row in tabela.rows[1:]:
            celulas = [c.text for c in row.cells]
            registro = {
                _ORIG_POR_NORM[campo]: (celulas[j].strip() if j < len(celulas) else "")
                for campo, j in idx_por_campo.items()
            }
            if not any(v.strip() for v in registro.values()):
                continue  # linha vazia
            if _cabecalho_casa(registro.values()):
                continue  # repetição do cabeçalho
            linhas.append(registro)

        if linhas:
            texto = "\n".join(p.text for p in doc.paragraphs)
            return ResultadoExtracao(
                df=pd.DataFrame(linhas),
                numero_psc=_detectar_psc(texto),
                paginas=0,
                metodo="docx",
            )

    raise ErroExtracao(
        "Tabela 'Lista de equipamentos / Material utilizado' não encontrada no DOCX. "
        "Se a tabela estiver como imagem, use a versão em PDF."
    )


# ---------------------------------------------------------------------------
# DOC (via Microsoft Word / win32com)
# ---------------------------------------------------------------------------

_WD_FORMAT_DOCX = 16  # wdFormatDocumentDefault (.docx)


def _word_para_docx(conteudo: bytes) -> bytes:
    """Converte bytes de um .DOC em bytes .DOCX usando o Microsoft Word."""
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    word = None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            origem = Path(tmp) / "entrada.doc"
            destino = Path(tmp) / "saida.docx"
            origem.write_bytes(conteudo)

            word = win32.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False
            documento = word.Documents.Open(str(origem), ReadOnly=True)
            documento.SaveAs(str(destino), FileFormat=_WD_FORMAT_DOCX)
            documento.Close(False)
            return destino.read_bytes()
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:  # noqa: BLE001
                pass
        pythoncom.CoUninitialize()


def extrair_tabela_doc(conteudo: bytes, nome_arquivo: str) -> ResultadoExtracao:
    """Extrai a tabela de um .DOC (converte para DOCX via Word e reusa o leitor)."""
    try:
        docx_bytes = _word_para_docx(conteudo)
    except Exception as exc:  # noqa: BLE001
        raise ErroExtracao(
            "Para ler arquivos .DOC é preciso o Microsoft Word instalado "
            f"(a conversão automática falhou: {exc}). "
            "Como alternativa, salve o arquivo como .DOCX ou PDF."
        ) from exc

    resultado = extrair_tabela_docx(docx_bytes, nome_arquivo)
    resultado.metodo = "doc"
    return resultado
